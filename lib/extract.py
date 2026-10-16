"""Text extraction for PULSE content index. Never raises; returns (text, state)."""
import os
import zipfile
import warnings

# Third-party readers are chatty about files they can still read perfectly
# well (openpyxl default styles, print areas; PyMuPDF's fitz alias). None of
# it is actionable here, and it drowns the progress line.
warnings.filterwarnings("ignore", category=UserWarning)
warnings.filterwarnings("ignore", category=DeprecationWarning)

MAX_CHARS = 180_000

TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml", ".xml",
            ".html", ".htm", ".py", ".js", ".ts", ".ps1", ".sh", ".sql",
            ".ini", ".cfg", ".toml", ".bat", ".cmd", ".jsonl", ".log"}
PDF_EXT   = {".pdf"}
XL_EXT    = {".xlsx", ".xlsm", ".xltx"}
DOC_EXT   = {".docx", ".dotx"}
PPT_EXT   = {".pptx", ".potx"}

# Per-type size caps in bytes
CAPS = {"pdf": 40 << 20, "xl": 30 << 20, "doc": 20 << 20,
        "ppt": 30 << 20, "text": 8 << 20}

SUPPORTED = TEXT_EXT | PDF_EXT | XL_EXT | DOC_EXT | PPT_EXT


def kind(ext):
    if ext in PDF_EXT:  return "pdf"
    if ext in XL_EXT:   return "xl"
    if ext in DOC_EXT:  return "doc"
    if ext in PPT_EXT:  return "ppt"
    if ext in TEXT_EXT: return "text"
    return None


def _pdf(path):
    try:
        import pymupdf as fitz          # current name
    except ImportError:
        import fitz                     # older installs
    parts = []
    with fitz.open(path) as d:
        for i, page in enumerate(d):
            if i >= 150:
                break
            parts.append(page.get_text("text"))
            if sum(map(len, parts)) > MAX_CHARS:
                break
    return "\n".join(parts)


def _xl(path):
    import openpyxl
    parts = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            parts.append("### " + str(ws.title))
            n = 0
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None]
                if vals:
                    parts.append("\t".join(vals))
                n += 1
                if n > 4000 or sum(map(len, parts)) > MAX_CHARS:
                    break
            if sum(map(len, parts)) > MAX_CHARS:
                break
    finally:
        wb.close()
    return "\n".join(parts)


def _doc(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _ppt(path):
    # Pull raw text from slide XML without python-pptx dependency
    import re
    parts = []
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist()
                 if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        for n in sorted(names)[:120]:
            xml = z.read(n).decode("utf-8", "ignore")
            parts.extend(re.findall(r"<a:t>(.*?)</a:t>", xml, re.S))
    return "\n".join(parts)


def _text(path):
    with open(path, "rb") as f:
        raw = f.read(CAPS["text"])
    for enc in ("utf-8", "utf-8-sig", "cp1256", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "ignore")


_FN = {"pdf": _pdf, "xl": _xl, "doc": _doc, "ppt": _ppt, "text": _text}


def extract(path, ext, size):
    """-> (text_or_None, state) where state 1=done 2=skip 3=error"""
    k = kind(ext)
    if not k:
        return None, 2
    if size <= 0 or size > CAPS[k]:
        return None, 2
    try:
        txt = _FN[k](path)
    except Exception:
        return None, 3
    if not txt or not txt.strip():
        return None, 2
    return txt[:MAX_CHARS], 1
